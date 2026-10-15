"""Parsing helpers: money, dates, CSV/XLSX rows, WeChat/Alipay statements.

Money is always parsed straight to integer fen via Decimal — never floats.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any
from zoneinfo import ZoneInfo

from .errors import HeliosError
from .filestore import decode_statement_bytes
from .store import doc_hash

SH_TZ = ZoneInfo("Asia/Shanghai")

# Guide § limits: 10,000 rows per import_sheet call.
MAX_IMPORT_ROWS = 10_000

DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y年%m月%d日 %H:%M:%S",
    "%Y年%m月%d日 %H:%M",
    "%Y年%m月%d日",
)

NUMBER_RE = re.compile(r"-?\d+(?:\.\d+)?")
AMOUNT_RE = re.compile(
    r"[¥￥]?\s*(\d{1,3}(?:,\d{3})*(?:\.\d+)?|\d+(?:\.\d+)?)\s*元?")
DATE_RE = re.compile(r"\d{4}[-/年]\d{1,2}[-/月]\d{1,2}日?")
DUE_KEYWORDS = ("到期", "截止", "最迟", "期限", "before", "due")
REF_RE = re.compile(
    r"(?:发票号|发票号码|Invoice|INV|No\.?|编号|单号)[:#：]?\s*([A-Za-z0-9\-/]+)",
    re.IGNORECASE)

DIRECTIONS_CN = {
    "收": "in", "收入": "in", "贷": "in", "in": "in",
    "支": "out", "支出": "out", "借": "out", "out": "out",
}
OBLIGATION_DIRECTIONS_CN = {
    "应收": "receivable", "receivable": "receivable", "收": "receivable",
    "收入": "receivable",
    "应付": "payable", "payable": "payable", "支": "payable", "支出": "payable",
}


def parse_money(text: Any) -> int | None:
    """Money text ('¥1,234.56', 128.5, '100元') → integer fen, or None."""
    if text is None:
        return None
    if isinstance(text, int):
        return text
    if isinstance(text, float):  # XLSX numeric cells
        text = f"{text:.2f}"
    s = str(text).strip().replace(",", "").replace("，", "")
    m = NUMBER_RE.search(s)
    if not m:
        return None
    try:
        fen = (Decimal(m.group(0)) * 100).quantize(Decimal("1"))
    except InvalidOperation:
        return None
    return int(fen)


def format_money(fen: int) -> str:
    sign = "-" if fen < 0 else ""
    a = abs(fen)
    return f"{sign}{a // 100}.{a % 100:02d}"


def parse_datetime(text: Any) -> datetime | None:
    s = str(text or "").strip()
    if not s:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt).replace(tzinfo=SH_TZ)
        except ValueError:
            continue
    try:
        dt = datetime.fromisoformat(s)
    except ValueError:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=SH_TZ)


def parse_amounts(text: str) -> list[int]:
    out = []
    for m in AMOUNT_RE.findall(text or ""):
        try:
            out.append(int((Decimal(m.replace(",", "")) * 100)
                           .quantize(Decimal("1"))))
        except InvalidOperation:
            continue
    return out


def find_dates(text: str) -> list[datetime]:
    out = []
    for m in DATE_RE.finditer(text or ""):
        dt = parse_datetime(m.group(0))
        if dt:
            out.append(dt)
    return out


def find_due_date(text: str) -> datetime | None:
    """Prefer a date on a line that mentions a due keyword; else first date."""
    for line in (text or "").splitlines():
        if any(k in line.lower() for k in DUE_KEYWORDS):
            dates = find_dates(line) or find_dates(text)
            if dates:
                return dates[0]
    dates = find_dates(text)
    return dates[0] if dates else None


def find_refs(text: str) -> list[str]:
    return REF_RE.findall(text or "")


# --------------------------------------------------------------- row readers

def _clean_rows(raw_rows) -> list[list[str]]:
    rows = []
    for r in raw_rows:
        vals = [str(c).strip() if c is not None else "" for c in r]
        if any(vals):
            rows.append(vals)
    return rows


def _grid_to_dicts(header: list[str], rows: list[list[str]]) -> list[dict]:
    out = []
    for vals in rows:
        padded = vals + [""] * (len(header) - len(vals))
        out.append(dict(zip(header, padded[:len(header)])))
    return out


def _check_row_limit(rows: list, filename: str = "") -> None:
    """Guide § limits: 10,000 rows per import_sheet call."""
    if len(rows) > MAX_IMPORT_ROWS:
        raise HeliosError(
            "validation_error",
            f"{len(rows)} data rows exceeds the {MAX_IMPORT_ROWS} row limit"
            + (f" for '{filename}'" if filename else "")
            + "; split the file and import in batches",
            details={"rows": len(rows), "limit": MAX_IMPORT_ROWS})


def rows_from_csv(content: bytes, filename: str = "") -> tuple[list[str], list[dict]]:
    text = decode_statement_bytes(content, filename)
    grid = _clean_rows(csv.reader(io.StringIO(text)))
    if not grid:
        return [], []
    _check_row_limit(grid[1:], filename)
    return grid[0], _grid_to_dicts(grid[0], grid[1:])


def rows_from_xlsx(content: bytes, filename: str = "") -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HeliosError(
            "unsupported_format",
            "file is not a readable XLSX workbook") from exc
    ws = wb.active
    grid = _clean_rows(ws.iter_rows(values_only=True))
    if not grid:
        return [], []
    _check_row_limit(grid[1:], filename)
    return grid[0], _grid_to_dicts(grid[0], grid[1:])


# ---------------------------------------------------------- statement parsing

def _find_statement_grid(grid: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    for i, row in enumerate(grid):
        if len(row) < 4:
            continue
        if any(("交易时间" in c) or ("交易日期" in c) or ("记账日期" in c)
               or c in ("日期", "时间") for c in row):
            return row, grid[i + 1:]
    raise HeliosError(
        "unsupported_format",
        "could not find a statement header row (expected 交易时间/交易日期…)")


def statement_grid_from_csv(content: bytes, filename: str = ""):
    text = decode_statement_bytes(content, filename)
    grid = _clean_rows(csv.reader(io.StringIO(text)))
    return _find_statement_grid(grid)


def statement_grid_from_xlsx(content: bytes):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise HeliosError(
            "unsupported_format",
            "file is not a readable XLSX workbook") from exc
    return _find_statement_grid(_clean_rows(wb.active.iter_rows(values_only=True)))


def _pick(d: dict, *names: str) -> str:
    for n in names:
        v = d.get(n)
        if v not in (None, "", "-"):
            return v
    return ""


def _norm_wechat(d: dict) -> dict | None:
    direction = DIRECTIONS_CN.get(str(d.get("收/支", "")).strip())
    if not direction:
        return None  # 不计收支 — not income or expense
    amount = parse_money(_pick(d, "金额(元)", "金额（元）", "金额"))
    ref = _pick(d, "交易单号")
    occurred = parse_datetime(_pick(d, "交易时间"))
    if amount is None or not ref or occurred is None:
        return None
    return {
        "occurred_at": occurred,
        "direction": direction,
        "amount": amount,
        "counterparty": _pick(d, "交易对方") or "未知",
        "channel_ref": ref,
        "needs_review": False,
        "meta": {
            k: v for k, v in {
                "type": _pick(d, "交易类型"),
                "product": _pick(d, "商品"),
                "method": _pick(d, "支付方式"),
                "status": _pick(d, "当前状态"),
                "remark": _pick(d, "备注"),
                "merchant_ref": _pick(d, "商户单号"),
            }.items() if v
        },
    }


def _norm_alipay(d: dict) -> dict | None:
    direction = DIRECTIONS_CN.get(str(d.get("收/支", "")).strip())
    if not direction:
        return None
    amount = parse_money(_pick(d, "金额"))
    ref = _pick(d, "交易流水号")
    occurred = parse_datetime(_pick(d, "交易时间"))
    if amount is None or not ref or occurred is None:
        return None
    from .filestore import mask_account
    meta = {
        k: v for k, v in {
            "category": _pick(d, "交易分类"),
            "account": mask_account(_pick(d, "对方账号")),
            "product": _pick(d, "商品说明"),
            "method": _pick(d, "收/付款方式"),
            "status": _pick(d, "交易状态"),
            "remark": _pick(d, "备注"),
            "merchant_ref": _pick(d, "商家订单号"),
        }.items() if v
    }
    return {
        "occurred_at": occurred,
        "direction": direction,
        "amount": amount,
        "counterparty": _pick(d, "交易对方") or "未知",
        "channel_ref": ref,
        "needs_review": False,
        "meta": meta,
    }


def _norm_generic(d: dict, header: list[str]) -> dict | None:
    occurred = parse_datetime(
        _pick(d, "交易时间", "交易日期", "记账日期", "日期", "时间"))
    if occurred is None:
        return None
    needs_review = False
    if "收入金额" in header and "支出金额" in header:
        a_in = parse_money(_pick(d, "收入金额"))
        a_out = parse_money(_pick(d, "支出金额"))
        if a_in:
            amount, direction = a_in, "in"
        elif a_out:
            amount, direction = a_out, "out"
        else:
            return None
    else:
        amount = parse_money(_pick(d, "金额", "交易金额", "发生额"))
        if amount is None:
            return None
        direction = DIRECTIONS_CN.get(
            _pick(d, "收/支", "借贷", "方向").lower().strip())
        if direction is None:
            direction, needs_review = "in", True
    ref = _pick(d, "交易单号", "流水号", "凭证号", "参考号", "单号")
    if not ref:
        ref = "row_" + doc_hash(json.dumps(d, ensure_ascii=False, sort_keys=True))
    return {
        "occurred_at": occurred,
        "direction": direction,
        "amount": amount,
        "counterparty": _pick(d, "交易对方", "对方户名", "户名", "对方",
                              "付款人", "收款人") or "未知",
        "channel_ref": ref,
        "needs_review": needs_review,
        "meta": {
            k: v for k, v in {
                "status": _pick(d, "状态", "交易状态"),
                "remark": _pick(d, "备注", "摘要"),
            }.items() if v
        },
    }


def detect_statement_channel(header: list[str], default_channel: str = "bank") -> str:
    """Identify which wallet/bank layout a statement header uses."""
    header = [h.strip() for h in header]
    if any("交易单号" in h for h in header):
        return "wechat"
    if any("交易流水号" in h for h in header):
        return "alipay"
    return default_channel


def normalize_statement_row(d: dict, header: list[str], channel: str) -> dict | None:
    """Normalize one statement row, or None if it is not a usable transaction.

    Row-by-row so the caller can report the exact file row it dropped.
    """
    if channel == "wechat":
        return _norm_wechat(d)
    if channel == "alipay":
        return _norm_alipay(d)
    return _norm_generic(d, header)


def parse_statement_rows(header: list[str], rows: list[list[str]],
                         default_channel: str = "bank") -> tuple[str, list[dict]]:
    """Normalize a statement grid into transaction items.

    Returns (channel, items); footer/summary rows are dropped. Use
    `normalize_statement_row` directly when you need to know which file row was
    dropped — this function returns only the rows that survived.
    """
    header = [h.strip() for h in header]
    channel = detect_statement_channel(header, default_channel)
    dicts = _grid_to_dicts(header, rows)
    items = [normalize_statement_row(d, header, channel) for d in dicts]
    return channel, [i for i in items if i]


def detect_channel(*names: str) -> str:
    for name in names:
        if not name:
            continue
        n = name.lower()
        if "微信" in name or "wechat" in n:
            return "wechat"
        if "支付宝" in name or "alipay" in n:
            return "alipay"
    return "bank"
